import asyncio
import json
import os
import threading
import time
import tkinter as tk
import unittest
import uuid
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock, patch

import pandas as pd
from openpyxl import Workbook, load_workbook

from app.config import AppConfig
from app.services.api_manager import APIManager
from app.services.data_manager import DataManager
from app.services.ollama_api_manager import OllamaAPIManager
from app.services.processing import ProcessingSettings
from app.ui.dispatcher import UIDispatcher
from app.ui.main_window import ExcelAIAssistantApp
from app.ui.dialogs.preferences_dialog import PreferencesDialog
from app.ui.dialogs.ollama_settings_dialog import OllamaSettingsDialog
from app.utils.theme_manager import ThemeManager


class WorkspaceFixture(unittest.TestCase):
    def setUp(self):
        credentials = patch.dict(os.environ, {'OPENAI_API_KEY': '', 'OPENROUTER_API_KEY': ''})
        credentials.start()
        self.addCleanup(credentials.stop)
        self.directory = Path(__file__).resolve().parent / ('fixtures-' + uuid.uuid4().hex)
        self.directory.mkdir()
        def cleanup():
            for item in self.directory.iterdir():
                item.unlink()
            self.directory.rmdir()
        self.addCleanup(cleanup)

    def csv_manager(self, content='id,name\n0012,NA\n'):
        source = self.directory / 'source.csv'
        source.write_text(content, encoding='utf-8')
        manager = DataManager()
        self.assertTrue(manager.load_file(source)[0])
        return source, manager


class SafetyTests(WorkspaceFixture):
    def test_openrouter_environment_key_and_preferences_keep_keys_separate_and_unsaved(self):
        path = self.directory / 'settings.json'
        with patch.object(AppConfig, '_get_config_path', return_value=path), \
                patch.dict(os.environ, {'OPENROUTER_API_KEY': 'router-environment-key'}):
            config = AppConfig()
        self.assertEqual(config.get('openrouter_api_key'), 'router-environment-key')
        config.set('api_key', 'openai-session-key')
        config.set_active_model('openrouter', 'example/custom-model')
        self.assertEqual(config.get_active_model(), {'api_type': 'openrouter', 'model': 'example/custom-model'})
        self.assertTrue(config.save())
        saved = json.loads(path.read_text())
        self.assertNotIn('api_key', saved)
        self.assertNotIn('openrouter_api_key', saved)
        self.assertNotIn('router-environment-key', path.read_text())
        self.assertEqual(saved['openrouter_model'], 'example/custom-model')

    def test_openrouter_uses_its_endpoint_and_selected_model_for_generation_and_test(self):
        with patch('app.services.api_manager.OpenAI') as factory:
            manager = APIManager(api_type='openrouter', api_key='router-test-key', model='example/custom-model')
        self.assertEqual(factory.call_args.kwargs['base_url'], 'https://openrouter.ai/api/v1')
        self.assertEqual(factory.call_args.kwargs['api_key'], 'router-test-key')
        self.assertEqual(factory.call_args.kwargs['max_retries'], 0)
        client = factory.return_value
        client.chat.completions.create.return_value = SimpleNamespace(choices=[
            SimpleNamespace(finish_reason='stop', message=SimpleNamespace(content='reviewed value'))])
        self.assertEqual(manager.process_single_cell('old', 'system', 'transform', context_data={'context': 'selected'}),
                         (True, 'reviewed value', None))
        request = client.chat.completions.create.call_args.kwargs
        self.assertEqual(request['model'], 'example/custom-model')
        self.assertIn('selected', request['messages'][1]['content'])
        self.assertTrue(manager.test_connection()[0])
        self.assertEqual(client.chat.completions.create.call_args.kwargs['model'], 'example/custom-model')

    def test_switching_cloud_providers_clears_the_previous_key_and_client(self):
        manager = APIManager()
        manager.client = Mock()
        manager.api_key = 'openai-only-key'
        manager.set_api_type('openrouter')
        self.assertIsNone(manager.client)
        self.assertEqual(manager.api_key, '')
        self.assertFalse(manager.initialize())
        with patch('app.services.api_manager.OpenAI') as factory:
            self.assertTrue(manager.initialize('router-only-key'))
        self.assertEqual(factory.call_args.kwargs['api_key'], 'router-only-key')
        manager.set_api_type('openai')
        self.assertIsNone(manager.client)
        self.assertEqual(manager.api_key, '')

    def test_openrouter_incomplete_and_error_responses_are_not_successful_results(self):
        manager = APIManager(api_type='openrouter', model='example/custom-model')
        manager.client = Mock()
        manager.client.chat.completions.create.return_value = SimpleNamespace(choices=[
            SimpleNamespace(finish_reason='length', message=SimpleNamespace(content='partial'))])
        self.assertFalse(manager.process_single_cell('old', '', '')[0])
        manager.client.chat.completions.create.side_effect = RuntimeError('Insufficient credits')
        success, result, error = manager.process_single_cell('old', '', '')
        self.assertFalse(success)
        self.assertIsNone(result)
        self.assertIn('Insufficient credits', error)

    def test_identifiers_and_literal_na_survive_csv_load(self):
        _, manager = self.csv_manager()
        self.assertEqual(manager.get_cell_value(0, 'id'), '0012')
        self.assertEqual(manager.get_cell_value(0, 'name'), 'NA')

    def test_external_change_blocks_save_and_keeps_edits(self):
        source, manager = self.csv_manager()
        manager.update_cell(0, 'name', 'edited')
        source.write_text('external version', encoding='utf-8')
        self.assertFalse(manager.save_file()[0])
        self.assertEqual(source.read_text(), 'external version')
        self.assertEqual(manager.get_cell_value(0, 'name'), 'edited')
        self.assertTrue(manager.modified)

    def test_existing_file_is_backed_up_before_save(self):
        source, manager = self.csv_manager()
        before = source.read_bytes()
        manager.update_cell(0, 'name', 'edited')
        self.assertTrue(manager.save_file()[0])
        self.assertEqual(Path(str(source) + '.bak').read_bytes(), before)

    def test_csv_neutralizes_formula_prefixes(self):
        source, manager = self.csv_manager()
        manager.update_cell(0, 'name', '=HYPERLINK("https://example.invalid")')
        self.assertTrue(manager.save_file()[0])
        result = pd.read_csv(source, dtype=str)
        self.assertTrue(result.loc[0, 'name'].startswith("'="))

    def test_excel_generated_formula_is_text(self):
        _, manager = self.csv_manager()
        manager.update_cell(0, 'name', '=1+2')
        target = self.directory / 'safe.xlsx'
        self.assertTrue(manager.save_file(target)[0])
        book = load_workbook(target)
        try:
            self.assertEqual(book.active['B2'].value, '=1+2')
            self.assertEqual(book.active['B2'].data_type, 's')
        finally:
            book.close()

    def test_undo_restores_original_excel_formula(self):
        source = self.directory / 'formulas.xlsx'
        book = Workbook()
        book.active.append(['formula'])
        book.active.append(['=1+2'])
        book.save(source)
        book.close()
        manager = DataManager()
        self.assertTrue(manager.load_file(source)[0])
        manager.update_cell(0, 'formula', 'changed')
        self.assertTrue(manager.undo())
        self.assertFalse(manager.modified)
        self.assertTrue(manager.save_file()[0])
        book = load_workbook(source)
        try:
            self.assertEqual(book.active['A2'].value, '=1+2')
        finally:
            book.close()

    def test_advanced_workbook_parts_survive_cell_edits(self):
        import zipfile
        source = self.directory / 'advanced.xlsx'
        book = Workbook()
        book.active.append(['name'])
        book.active.append(['old'])
        book.save(source)
        book.close()
        with zipfile.ZipFile(source, 'a') as archive:
            archive.writestr('xl/slicers/slicer1.xml', '<unsupported/>')
        with zipfile.ZipFile(source) as archive:
            before = {name: archive.read(name) for name in archive.namelist()}
        manager = DataManager()
        self.assertTrue(manager.load_file(source)[0])
        manager.update_cell(0, 'name', 'new')
        success, error = manager.save_file()
        self.assertTrue(success, error)
        with zipfile.ZipFile(source) as archive:
            for name, payload in before.items():
                if name not in ('xl/worksheets/sheet1.xml', 'xl/workbook.xml'):
                    self.assertEqual(archive.read(name), payload)
        self.assertFalse(manager.modified)

    def test_failed_result_cannot_overwrite_cell(self):
        _, manager = self.csv_manager()
        applied, failed = manager.update_range([{'row': 0, 'col': 'name', 'result': 'bad', 'success': False}])
        self.assertEqual((applied, failed), (0, 1))
        self.assertEqual(manager.get_cell_value(0, 'name'), 'NA')

    def test_config_does_not_persist_api_key(self):
        path = self.directory / 'settings.json'
        with patch.object(AppConfig, '_get_config_path', return_value=path):
            config = AppConfig()
        config.set('api_key', 'test-secret-do-not-persist')
        self.assertTrue(config.save())
        self.assertNotIn('api_key', json.loads(path.read_text()))
        self.assertNotIn('test-secret-do-not-persist', path.read_text())

    def test_bad_config_types_fall_back_to_valid_defaults(self):
        path = self.directory / 'settings.json'
        path.write_text(json.dumps({'api_type': None, 'recent_files': {}, 'batch_size': -5,
                                    'temperature': 'bad', 'ollama_url': None, 'prompts': []}))
        with patch.object(AppConfig, '_get_config_path', return_value=path):
            config = AppConfig()
        self.assertEqual(config.get('api_type'), 'openai')
        self.assertEqual(config.get('batch_size'), 10)
        self.assertEqual(config.get('recent_files'), [])
        self.assertIsInstance(config.get_prompt_templates(), dict)

    def test_default_reset_and_malformed_settings_keep_logging_opt_in(self):
        path = self.directory / 'settings.json'
        path.write_text(json.dumps({'save_logs': 'invalid'}))
        with patch.object(AppConfig, '_get_config_path', return_value=path):
            config = AppConfig()
        self.assertFalse(config.get('save_logs'))
        config.set('save_logs', True)
        config.restore_defaults()
        self.assertFalse(config.get('save_logs'))

    def test_invalid_settings_and_ollama_urls_are_rejected(self):
        for values in [(0, .3, 150), (10, float('nan'), 150), (10, .3, -1)]:
            with self.assertRaises(ValueError):
                ProcessingSettings.parse(*values)
        for url in ['file:///tmp/data', 'http://user:pass@localhost:11434', 'https://localhost?token=secret']:
            with self.assertRaises(ValueError):
                OllamaAPIManager(base_url=url)

    def test_clearing_key_invalidates_old_client(self):
        manager = APIManager()
        manager.api_key = 'old-test-key'
        manager.client = Mock()
        self.assertFalse(manager.initialize(''))
        self.assertIsNone(manager.client)
        self.assertEqual(manager.api_key, '')

    def test_truncated_openai_output_is_rejected(self):
        manager = APIManager()
        manager.client = Mock()
        manager.client.chat.completions.create.return_value = SimpleNamespace(choices=[
            SimpleNamespace(finish_reason='length', message=SimpleNamespace(content='partial'))])
        success, result, error = manager.process_single_cell('old', '', '')
        self.assertFalse(success)
        self.assertIsNone(result)

    def test_rate_window_reports_wait_instead_of_discarding_rest_of_job(self):
        manager = APIManager()
        manager.request_count = manager.max_requests_per_minute
        manager.request_start_time = time.time()
        self.assertGreater(manager.rate_limit_delay(), 59)


class DesktopTests(WorkspaceFixture):
    def setUp(self):
        super().setUp()
        try:
            self.root = tk.Tk()
        except tk.TclError as error:
            self.skipTest(f'Tk unavailable: {error}')
        self.root.withdraw()
        config_path = self.directory / 'settings.json'
        with patch.object(AppConfig, '_get_config_path', return_value=config_path):
            self.config = AppConfig()
        self.config.set('save_logs', False)
        self.config.set('api_key', '')
        theme = ThemeManager(self.root)
        theme.set_theme('light')
        self.app = ExcelAIAssistantApp(self.root, self.config, theme)
        self.addCleanup(self.close_app)

    def close_app(self):
        self.app.dispatcher.close()
        self.app.batch_processor.dispatcher.close()
        self.app.loop.call_soon_threadsafe(self.app.loop.stop)
        self.app.loop_thread.join(timeout=3)
        self.app.loop.close()
        self.root.destroy()

    def test_preferences_opens_and_applies_settings(self):
        dialog = PreferencesDialog(self.root, self.config, self.app.api_manager, self.app._preferences_updated)
        dialog.batch_size_var.set('5')
        dialog.max_tokens_var.set('200')
        self.assertTrue(dialog._apply_settings())
        self.assertEqual(self.app.batch_size_var.get(), '5')
        self.assertEqual(self.app.max_tokens_var.get(), '200')
        self.assertFalse(self.app.save_after_batch_var.get())
        dialog.dialog.destroy()

    def test_openrouter_toolbar_keeps_credentials_and_custom_models_separate(self):
        self.app.api_key_entry.insert(0, 'openai-test-key')
        self.app.model_var.set('openai-custom')
        self.app.openrouter_api_key_entry.insert(0, 'router-test-key')
        with patch('app.services.api_manager.OpenAI') as factory:
            self.app.api_type_var.set('openrouter')
            self.app._on_api_type_changed()
            self.app.model_var.set('example/custom-model')
            manager = self.app._configured_manager()
            self.assertEqual(factory.call_args.kwargs['api_key'], 'router-test-key')
            self.assertEqual(manager.api_type.value, 'openrouter')
            self.assertEqual(manager.model, 'example/custom-model')
            self.assertEqual(self.app.openrouter_frame.winfo_manager(), 'pack')
            self.assertEqual(self.app.openai_frame.winfo_manager(), '')
            self.app.api_type_var.set('openai')
            self.app._on_api_type_changed()
            self.assertEqual(self.app.model_var.get(), 'openai-custom')
            self.assertEqual(self.app.api_key_entry.get(), 'openai-test-key')
            self.assertEqual(factory.call_args.kwargs['api_key'], 'openai-test-key')
            self.assertNotIn('base_url', factory.call_args.kwargs)
            self.app.api_type_var.set('openrouter')
            self.app._on_api_type_changed()
        self.assertEqual(self.app.model_var.get(), 'example/custom-model')
        self.app._save_config()
        saved = self.config._config_file.read_text()
        self.assertNotIn('router-test-key', saved)
        self.assertNotIn('openai-test-key', saved)

    def test_openrouter_requires_its_own_key_and_namespaced_model(self):
        self.app.api_type_var.set('openrouter')
        self.app._on_api_type_changed()
        self.app.api_key_entry.insert(0, 'openai-key-is-not-a-router-key')
        with self.assertRaisesRegex(ValueError, 'OpenRouter API key'):
            self.app._configured_manager()
        self.app.openrouter_api_key_entry.insert(0, 'router-test-key')
        self.app.model_var.set('invalid-model')
        with self.assertRaisesRegex(ValueError, 'provider/model'):
            self.app._configured_manager()

    def test_openrouter_preferences_apply_and_transfer_confirmation_names_destination(self):
        dialog = PreferencesDialog(self.root, self.config, self.app.api_manager, self.app._preferences_updated)
        try:
            dialog.api_type_var.set('openrouter')
            dialog._api_type_changed()
            self.assertEqual(dialog.api_notebook.index('current'), 2)
            dialog.openrouter_api_key_var.set('router-preferences-key')
            dialog.openrouter_model_var.set('example/preferences-model')
            with patch('app.services.api_manager.OpenAI'):
                self.assertTrue(dialog._apply_settings())
                manager = self.app._configured_manager()
            self.assertEqual(manager.api_key, 'router-preferences-key')
            self.assertEqual(manager.model, 'example/preferences-model')
            self.assertNotIn('router-preferences-key', self.config._config_file.read_text())
            with patch('app.ui.main_window.messagebox.askyesno', return_value=True) as confirm:
                self.assertTrue(self.app._confirm_transfer(manager, 1, ['name'], []))
            text = confirm.call_args.args[1]
            self.assertIn('OpenRouter', text)
            self.assertIn('upstream model provider', text)
            self.assertNotIn('router-preferences-key', text)
        finally:
            dialog.dialog.destroy()

    def test_usage_estimate_rejects_invalid_token_settings(self):
        self.app.data_manager.df = pd.DataFrame({'name': ['sample']})
        self.app.max_tokens_var.set('invalid')
        with patch.object(self.app, '_get_selected_range', return_value=((0, 1), ['name'])), \
                patch('app.ui.main_window.messagebox.showerror') as alert:
            self.app._calculate_cost()
        alert.assert_called_once()

    def test_worksheet_selection_loads_requested_sheet_and_preserves_unsaved_edits(self):
        source = self.directory / 'sheets.xlsx'
        book = Workbook()
        book.active.title = 'First'
        book.active.append(['name'])
        book.active.append(['first'])
        second = book.create_sheet('Review, 2026')
        second.append(['name'])
        second.append(['second'])
        book.save(source)
        book.close()
        self.assertTrue(self.app.data_manager.load_file(str(source))[0])
        self.app._update_after_file_load(True, '', str(source))
        self.assertEqual(tuple(self.app.sheet_combobox['values']), ('First', 'Review, 2026'))
        self.app.sheet_var.set('Review, 2026')
        self.app._on_sheet_selected()
        deadline = time.monotonic() + 5
        while self.app._operation and time.monotonic() < deadline:
            self.root.update()
            time.sleep(.005)
        self.assertIsNone(self.app._operation)
        self.assertEqual(self.app.data_manager.sheet_name, 'Review, 2026')
        self.assertEqual(self.app.data_manager.get_cell_value(0, 'name'), 'second')
        self.app.data_manager.update_cell(0, 'name', 'unsaved')
        self.app.sheet_var.set('First')
        with patch('app.ui.main_window.messagebox.askyesno', return_value=False):
            self.app._on_sheet_selected()
        self.assertEqual(self.app.data_manager.sheet_name, 'Review, 2026')
        self.assertEqual(self.app.sheet_var.get(), 'Review, 2026')
        self.assertEqual(self.app.data_manager.get_cell_value(0, 'name'), 'unsaved')

    def test_dispatcher_delivers_on_ui_thread_and_ignores_closed_posts(self):
        dispatcher = UIDispatcher(self.root)
        calls = []
        ui_thread = threading.get_ident()
        worker = threading.Thread(target=lambda: dispatcher.post(lambda: calls.append(threading.get_ident())))
        worker.start()
        worker.join()
        self.root.after_cancel(dispatcher.timer)
        dispatcher._drain()
        self.assertEqual(calls, [ui_thread])
        dispatcher.close()
        dispatcher.post(lambda: calls.append(-1))
        self.assertEqual(calls, [ui_thread])

    def test_ollama_settings_does_not_switch_live_provider(self):
        with patch.object(APIManager, 'get_available_models', return_value=[]):
            dialog = OllamaSettingsDialog(self.root, self.config, self.app.api_manager)
            deadline = time.monotonic() + .15
            while time.monotonic() < deadline:
                self.root.update()
                time.sleep(.005)
            self.assertEqual(self.app.api_manager.api_type.value, 'openai')
            dialog.dialog.destroy()

    def test_range_uses_visible_row_numbers(self):
        self.app.data_manager.df = pd.DataFrame({'name': ['a', 'b', 'c']})
        self.app.start_row.delete(0, tk.END)
        self.app.start_row.insert(0, '2')
        self.app.end_row.delete(0, tk.END)
        self.app.end_row.insert(0, '3')
        self.app.columns_entry.insert(0, 'name')
        self.assertEqual(self.app._get_selected_range(), ((1, 3), ['name']))

    def test_columns_with_commas_are_selected_as_single_columns(self):
        self.app.data_manager.df = pd.DataFrame({'Last, First': ['Smith, Alice']})
        self.app._update_column_lists()
        self.app.end_row.delete(0, tk.END)
        self.app.end_row.insert(0, '1')
        self.assertEqual(self.app._get_selected_range(), ((0, 1), ['Last, First']))

    def test_busy_document_rejects_cell_edits(self):
        self.app.data_manager.df = pd.DataFrame({'name': ['old']})
        self.app._operation = 'saving a file'
        with patch('app.ui.main_window.messagebox.showinfo'):
            self.assertFalse(self.app._cell_edited(0, 'name', 'new'))
        self.assertEqual(self.app.data_manager.get_cell_value(0, 'name'), 'old')
        self.app._operation = None

    def test_review_selection_applies_only_selected_proposal(self):
        self.app.data_manager.df = pd.DataFrame({'name': ['old1', 'old2']})
        self.app._run_document_id = self.app.data_manager.document_id
        self.app._run_autosave = False
        self.app._processing_completed([
            {'row': 0, 'col': 'name', 'original': 'old1', 'result': 'new1', 'success': True},
            {'row': 1, 'col': 'name', 'original': 'old2', 'result': 'new2', 'success': True}], 2, 0)
        self.assertEqual(self.app.data_manager.get_cell_value(0, 'name'), 'old1')
        dialog = next(widget for widget in self.root.winfo_children() if isinstance(widget, tk.Toplevel))
        tree = next(widget for widget in dialog.winfo_children() if isinstance(widget, tk.ttk.Treeview))
        tree.selection_set('1')
        def buttons(widget):
            for child in widget.winfo_children():
                if isinstance(child, tk.ttk.Button):
                    yield child
                yield from buttons(child)
        next(button for button in buttons(dialog) if button.cget('text') == 'Apply Selected').invoke()
        self.assertEqual(self.app.data_manager.get_cell_value(0, 'name'), 'old1')
        self.assertEqual(self.app.data_manager.get_cell_value(1, 'name'), 'new2')
        self.assertTrue(self.app.data_manager.modified)
        self.assertIsNone(self.app._operation)

    def test_batch_runs_through_dispatcher_and_waits_for_review(self):
        self.app.data_manager.df = pd.DataFrame({'name': ['old1', 'old2']})
        self.app.columns_entry.insert(0, 'name')
        self.app.end_row.delete(0, tk.END)
        self.app.end_row.insert(0, '2')
        manager = Mock()
        manager.rate_limit_delay.return_value = 0
        manager.process_single_cell.return_value = (True, 'proposed', None)
        with patch.object(self.app, '_select_context_columns', return_value=[]), patch.object(
                self.app, '_configured_manager', return_value=manager), patch.object(
                self.app, '_confirm_transfer', return_value=True):
            self.app._run_transformation()
        deadline = time.monotonic() + 5
        while self.app._operation != 'reviewing results' and time.monotonic() < deadline:
            self.root.update()
            time.sleep(.005)
        self.assertEqual(self.app._operation, 'reviewing results')
        self.assertEqual(len(self.app._pending_results), 2)
        self.assertEqual(self.app.data_manager.get_cell_value(0, 'name'), 'old1')
        self.assertFalse(self.app.data_manager.modified)
        self.app._apply_reviewed_results([self.app._pending_results[0]])
        self.assertEqual(self.app.data_manager.get_cell_value(0, 'name'), 'proposed')
        self.assertEqual(self.app.data_manager.get_cell_value(1, 'name'), 'old2')

    def test_closing_review_discards_all_proposals(self):
        self.app.data_manager.df = pd.DataFrame({'name': ['old']})
        self.app._run_document_id = self.app.data_manager.document_id
        self.app._run_autosave = False
        self.app._processing_completed([{'row': 0, 'col': 'name', 'original': 'old',
                                         'result': 'new', 'success': True}], 1, 0)
        dialog = next(widget for widget in self.root.winfo_children() if isinstance(widget, tk.Toplevel))
        self.root.tk.call(dialog.protocol('WM_DELETE_WINDOW'))
        self.assertEqual(self.app.data_manager.get_cell_value(0, 'name'), 'old')
        self.assertEqual(self.app._pending_results, [])
        self.assertIsNone(self.app._operation)

    def test_stale_preview_cannot_edit_another_document(self):
        self.app.data_manager.df = pd.DataFrame({'name': ['old']})
        with patch('app.ui.main_window.messagebox.showerror') as alert:
            self.app._apply_preview(0, 'name', 'wrong', Mock(), 'different-document', 'old')
        self.assertEqual(self.app.data_manager.get_cell_value(0, 'name'), 'old')
        alert.assert_called_once()


class DependencyTests(unittest.TestCase):
    def test_dependency_manifests_match_lock_metadata(self):
        import tomllib
        root = Path(__file__).resolve().parent.parent
        project = tomllib.loads((root / 'pyproject.toml').read_text())['project']
        requirements = [line for line in (root / 'requirements.txt').read_text().splitlines() if line and not line.startswith('#')]
        self.assertEqual(sorted(project['dependencies']), sorted(requirements))
        lock = tomllib.loads((root / 'uv.lock').read_text())
        package = next(package for package in lock['package'] if package['name'] == project['name'])
        from app import __version__
        self.assertEqual(project['version'], __version__)
        self.assertEqual(package['version'], __version__)
        metadata = [dependency['name'] + dependency['specifier'] for dependency in package['metadata']['requires-dist']]
        self.assertEqual(sorted(project['dependencies']), sorted(metadata))
        names = {package['name'] for package in lock['package']}
        for package in lock['package']:
            for dependency in package.get('dependencies', []):
                self.assertIn(dependency['name'], names)


if __name__ == '__main__':
    unittest.main()
