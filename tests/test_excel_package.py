"""Workbook fidelity and safety at the actual OOXML package boundary."""

from datetime import date
from pathlib import Path
import unittest
import uuid
import warnings
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook

from app.services.data_manager import DataManager
from app.services.excel_package import MAIN, WorkbookEditError, parse_xml, save_patched_workbook
from workbook_fixtures import create_representative_workbook, rewrite_parts


class PackageTests(unittest.TestCase):
    def setUp(self):
        self.directory = Path(__file__).resolve().parent / ("fixtures-" + uuid.uuid4().hex)
        self.directory.mkdir()
        self.addCleanup(self.cleanup)
        self.source = create_representative_workbook(self.directory)

    def cleanup(self):
        for path in self.directory.iterdir():
            path.unlink()
        self.directory.rmdir()

    @staticmethod
    def parts(path):
        with ZipFile(path) as archive:
            return {name: archive.read(name) for name in archive.namelist()}

    def manager(self, sheet="Sales"):
        manager = DataManager()
        with warnings.catch_warnings():
            warnings.simplefilter("ignore", UserWarning)
            self.assertTrue(manager.load_file(self.source, sheet_name=sheet)[0])
        return manager

    def test_real_chart_image_table_style_and_other_sheet_parts_stay_identical(self):
        before = self.parts(self.source)
        manager = self.manager()
        self.assertEqual(manager.get_cell_value(0, 'SKU'), '0012')
        self.assertTrue(manager.update_cell(0, "Description", "Changed & <safe>"))
        self.assertTrue(manager.save_file()[0])
        after = self.parts(self.source)
        self.assertEqual(set(before), set(after))
        changed = {name for name in before if before[name] != after[name]}
        self.assertEqual(changed, {"xl/worksheets/sheet1.xml", "xl/workbook.xml"})
        self.assertTrue(any(name.startswith("xl/charts/") for name in before))
        self.assertTrue(any(name.startswith("xl/media/") for name in before))
        book = load_workbook(self.source)
        try:
            sheet = book["Sales"]
            self.assertEqual(sheet["B2"].value, "Changed & <safe>")
            self.assertTrue(sheet["B2"].font.bold)
            self.assertEqual(sheet["B2"].comment.text, "Keep this comment and its VML drawing")
            self.assertEqual(sheet["E2"].value, "=C2*2")
            self.assertEqual(len(sheet._charts), 1)
            self.assertEqual(len(sheet._images), 1)
            self.assertIn("SalesData", sheet.tables)
            self.assertEqual(sheet.freeze_panes, "C2")
            self.assertIn("H2:I2", sheet.merged_cells)
            self.assertEqual(book["Hidden support"].sheet_state, "hidden")
            self.assertIn("Amounts", book.defined_names)
            self.assertTrue(book.calculation.fullCalcOnLoad)
        finally:
            book.close()

    def test_selecting_second_visible_sheet_edits_correct_relationship_part(self):
        before = self.parts(self.source)
        manager = self.manager("Review, 2026")
        self.assertEqual(manager.sheet_name, "Review, 2026")
        self.assertTrue(manager.update_cell(0, "Status", "approved"))
        self.assertTrue(manager.save_file()[0])
        after = self.parts(self.source)
        self.assertEqual(before["xl/worksheets/sheet1.xml"], after["xl/worksheets/sheet1.xml"])
        book = load_workbook(self.source, read_only=True)
        try:
            self.assertEqual(book["Review, 2026"]["B2"].value, "approved")
        finally:
            book.close()

    def test_opaque_advanced_parts_and_extension_prefixes_survive(self):
        before = self.parts(self.source)
        xml = before["xl/worksheets/sheet1.xml"].replace(b"<worksheet ",
            b'<worksheet xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
            b'xmlns:x14="http://schemas.microsoft.com/office/spreadsheetml/2009/9/main" mc:Ignorable="x14" ')
        extension = b'<extLst><ext uri="fixture"><x14:opaque marker="keep"/></ext></extLst>'
        xml = xml.replace(b"</worksheet>", extension + b"</worksheet>")
        opaque = {"xl/slicers/slicer-fixture.xml": b"<opaque>Slicer payload</opaque>",
                  "xl/pivotCache/pivotCacheRecords99.xml": b"<opaque>Cached pivot payload</opaque>",
                  "xl/activeX/activeX99.bin": b"\x00control\xff",
                  "xl/embeddings/oleObject99.bin": b"\x00embedded\xff",
                  "customXml/item99.xml": b"<opaque>Custom data</opaque>"}
        rewrite_parts(self.source, {"xl/worksheets/sheet1.xml": xml}, opaque)
        manager = self.manager()
        manager.update_cell(0, "Description", "updated")
        self.assertTrue(manager.save_file()[0])
        after = self.parts(self.source)
        for name, payload in opaque.items():
            self.assertEqual(after[name], payload)
        self.assertIn(extension, after["xl/worksheets/sheet1.xml"])
        self.assertIn(b'mc:Ignorable="x14"', after["xl/worksheets/sheet1.xml"])

    def test_macro_binary_and_vba_signature_payloads_are_copied_without_execution(self):
        # Synthetic package payloads verify preservation, not VBA runtime validity.
        source = self.directory / "macro-fixture.xlsm"
        source.write_bytes(self.source.read_bytes())
        parts = self.parts(source)
        types = parts["[Content_Types].xml"].replace(
            b"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml",
            b"application/vnd.ms-excel.sheet.macroEnabled.main+xml")
        opaque = {"xl/vbaProject.bin": b"synthetic-vba-project\x00\xff",
                  "xl/vbaProjectSignature.bin": b"synthetic-project-signature\x00\xff"}
        rewrite_parts(source, {"[Content_Types].xml": types}, opaque)
        target = self.directory / "macro-edited.xlsm"
        save_patched_workbook(source, target, "Sales", {"B2": "macro-safe data edit"})
        before, after = self.parts(source), self.parts(target)
        for name in before.keys() - {"xl/workbook.xml", "xl/worksheets/sheet1.xml"}:
            self.assertEqual(before[name], after[name])

    def test_document_signature_is_not_silently_invalidated(self):
        rewrite_parts(self.source, additions={"_xmlsignatures/sig1.xml": b"<signature/>"})
        manager = self.manager()
        before = self.source.read_bytes()
        manager.update_cell(0, "Description", "changed")
        success, error = manager.save_file()
        self.assertFalse(success)
        self.assertIn("document signature", error)
        self.assertEqual(self.source.read_bytes(), before)
        self.assertTrue(manager.modified)

    def test_protection_and_merged_child_edits_are_refused(self):
        with self.assertRaisesRegex(WorkbookEditError, "merged"):
            save_patched_workbook(self.source, self.directory / "merged.xlsx", "Sales", {"I2": "wrong"})
        parts = self.parts(self.source)
        xml = parts["xl/worksheets/sheet1.xml"].replace(b"</sheetData>", b'</sheetData><sheetProtection sheet="1"/>')
        rewrite_parts(self.source, {"xl/worksheets/sheet1.xml": xml})
        with self.assertRaisesRegex(WorkbookEditError, "protected"):
            save_patched_workbook(self.source, self.directory / "protected.xlsx", "Sales", {"B2": "wrong"})

    def test_array_shared_and_spill_ranges_are_preserved_and_not_overwritten(self):
        parts = self.parts(self.source)
        xml = parts["xl/worksheets/sheet1.xml"].replace(b"<f>C2*2</f>", b'<f t="array" ref="E2:F3">C2*2</f>')
        rewrite_parts(self.source, {"xl/worksheets/sheet1.xml": xml})
        with self.assertRaisesRegex(WorkbookEditError, "formula"):
            save_patched_workbook(self.source, self.directory / "array.xlsx", "Sales", {"F3": "wrong"})
        output = self.directory / "inputs.xlsx"
        save_patched_workbook(self.source, output, "Sales", {"B2": "input edit"})
        self.assertIn(b'<f t="array" ref="E2:F3">C2*2</f>', self.parts(output)["xl/worksheets/sheet1.xml"])

    def test_missing_blank_cells_and_blank_rows_keep_original_coordinates(self):
        book = Workbook()
        sheet = book.active
        sheet.append(["A", "B", "C"])
        sheet.append(["first", None, 1])
        sheet.append([None, None, None])
        sheet.append(["last", None, 2])
        source = self.directory / "blanks.xlsx"
        book.save(source)
        book.close()
        target = self.directory / "filled.xlsx"
        save_patched_workbook(source, target, "Sheet", {"B2": "filled", "A3": "new", "C3": 3})
        book = load_workbook(target)
        try:
            self.assertEqual(book.active["B2"].value, "filled")
            self.assertEqual(book.active["A3"].value, "new")
            self.assertEqual(book.active["C3"].value, 3)
            self.assertEqual(book.active["A4"].value, "last")
        finally:
            book.close()

    def test_stale_calc_chain_is_removed_with_its_relationship_and_content_type(self):
        parts = self.parts(self.source)
        relation = (b'<Relationship Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/calcChain" '
                    b'Target="calcChain.xml" Id="rIdCalc"/>')
        content_type = (b'<Override PartName="/xl/calcChain.xml" '
                        b'ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.calcChain+xml"/>')
        rewrite_parts(self.source, {
            "xl/_rels/workbook.xml.rels": parts["xl/_rels/workbook.xml.rels"].replace(b"</Relationships>", relation + b"</Relationships>"),
            "[Content_Types].xml": parts["[Content_Types].xml"].replace(b"</Types>", content_type + b"</Types>")},
            {"xl/calcChain.xml": f'<calcChain xmlns="{MAIN}"><c r="E2" i="1"/></calcChain>'.encode()})
        target = self.directory / "chain.xlsx"
        save_patched_workbook(self.source, target, "Sales", {"E2": "replaced formula"})
        result = self.parts(target)
        self.assertNotIn("xl/calcChain.xml", result)
        self.assertNotIn(b"calcChain", result["xl/_rels/workbook.xml.rels"])
        self.assertNotIn(b"calcChain", result["[Content_Types].xml"])

    def test_text_output_remains_text_and_handles_whitespace_unicode_and_xml(self):
        target = self.directory / "text.xlsx"
        value = '  =SUM(C2:C3) & <safe> "quotes" \u2603\r\nsecond line\rlast  '
        save_patched_workbook(self.source, target, "Sales", {"B2": value, "C3": 12.25, "D2": date(2026, 5, 1)})
        book = load_workbook(target)
        try:
            self.assertEqual(book["Sales"]["B2"].value, value)
            self.assertEqual(book["Sales"]["B2"].data_type, "s")
            self.assertEqual(book["Sales"]["C3"].value, 12.25)
            self.assertEqual(book["Sales"]["D2"].value, date(2026, 5, 1))
        finally:
            book.close()

    def test_illegal_text_and_entities_cannot_produce_invalid_packages(self):
        for value in ("bad\x00text", "x" * 32768):
            with self.assertRaises(WorkbookEditError):
                save_patched_workbook(self.source, self.directory / "invalid.xlsx", "Sales", {"B2": value})
        with self.assertRaises(WorkbookEditError):
            parse_xml(b'<!DOCTYPE x [<!ENTITY unsafe "expanded">]><x>&unsafe;</x>')

    def test_shared_formula_master_and_follower_are_both_guarded(self):
        parts = self.parts(self.source)
        xml = parts['xl/worksheets/sheet1.xml'].replace(
            b'<f>C2*2</f>', b'<f t="shared" si="0" ref="E2:E3">C2*2</f>').replace(
            b'<f>SUM(C2:C3)</f>', b'<f t="shared" si="0"/>')
        rewrite_parts(self.source, {'xl/worksheets/sheet1.xml': xml})
        for coordinate in ('E2', 'E3'):
            with self.assertRaisesRegex(WorkbookEditError, 'formula'):
                save_patched_workbook(self.source, self.directory / 'shared.xlsx', 'Sales', {coordinate: 'wrong'})

    def test_pivot_output_is_guarded_but_source_cells_remain_editable(self):
        parts = self.parts(self.source)
        rels = parts['xl/worksheets/_rels/sheet1.xml.rels']
        relation = (b'<Relationship Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/pivotTable" '
                    b'Target="../pivotTables/pivotTable99.xml" Id="rIdPivot"/>')
        rewrite_parts(self.source, {'xl/worksheets/_rels/sheet1.xml.rels': rels.replace(
            b'</Relationships>', relation + b'</Relationships>')},
            {'xl/pivotTables/pivotTable99.xml': f'<pivotTableDefinition xmlns="{MAIN}"><location ref="H10:I14"/></pivotTableDefinition>'.encode()})
        with self.assertRaisesRegex(WorkbookEditError, 'Pivot'):
            save_patched_workbook(self.source, self.directory / 'pivot.xlsx', 'Sales', {'I11': 'wrong'})
        output = self.directory / 'pivot-source.xlsx'
        save_patched_workbook(self.source, output, 'Sales', {'B2': 'source edit'})
        self.assertEqual(self.parts(output)['xl/pivotTables/pivotTable99.xml'],
                         self.parts(self.source)['xl/pivotTables/pivotTable99.xml'])

    def test_no_changes_preserves_every_package_payload(self):
        output = self.directory / 'unchanged.xlsx'
        save_patched_workbook(self.source, output, 'Sales', {})
        self.assertEqual(self.parts(output), self.parts(self.source))


if __name__ == "__main__":
    unittest.main()
