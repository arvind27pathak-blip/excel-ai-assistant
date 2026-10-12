"""Generate representative, non-sensitive workbook fixtures reproducibly."""

from datetime import datetime
from pathlib import Path
from zipfile import ZipFile

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.drawing.image import Image
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Font, PatternFill
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from PIL import Image as PillowImage


def rewrite_parts(path, replacements=None, additions=None):
    replacements, additions = replacements or {}, additions or {}
    with ZipFile(path) as archive:
        entries = [(info, archive.read(info.filename)) for info in archive.infolist()]
        comment = archive.comment
    with ZipFile(path, "w") as archive:
        archive.comment = comment
        for info, data in entries:
            archive.writestr(info, replacements.get(info.filename, data))
        for name, data in additions.items():
            archive.writestr(name, data)


def create_representative_workbook(directory):
    directory = Path(directory)
    directory.mkdir(parents=True, exist_ok=True)
    image_path = directory / "fixture-logo.png"
    PillowImage.new("RGB", (24, 24), color="#2266aa").save(image_path)
    book = Workbook()
    sheet = book.active
    sheet.title = "Sales"
    sheet.append(["SKU", "Description", "Amount", "AsOf", "Formula", "Blank"])
    sheet.append(["0012", "Original & <tag>", 12.5, datetime(2026, 1, 15), "=C2*2", None])
    sheet.append(["0023", "Another", 20, datetime(2026, 2, 16), "=SUM(C2:C3)", None])
    sheet.append(["0045", "  Keep spaces  ", 35, datetime(2026, 3, 17), "=C4+1", None])
    sheet["C2"].number_format = '#,##0.00'
    sheet["D2"].number_format = "yyyy-mm-dd"
    sheet["B2"].font = Font(name="Calibri", bold=True, color="2255AA")
    sheet["B2"].fill = PatternFill("solid", fgColor="FFF1AA")
    sheet["B2"].comment = Comment("Keep this comment and its VML drawing", "Fixture")
    sheet["A2"].hyperlink = "https://example.invalid/item/0012"
    sheet.row_dimensions[2].height = 28
    sheet.column_dimensions["B"].width = 34
    sheet.freeze_panes = "C2"
    sheet.print_title_rows = "1:1"
    sheet.print_area = "A1:I26"
    sheet.auto_filter.ref = "A1:F4"
    sheet.merge_cells("H2:I2")
    sheet["H2"] = "Merged note"
    table = Table(displayName="SalesData", ref="A1:F4")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium9", showRowStripes=True)
    sheet.add_table(table)
    validation = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0")
    validation.add("C2:C4")
    sheet.add_data_validation(validation)
    sheet.conditional_formatting.add("C2:C4", CellIsRule(operator="greaterThan", formula=["15"],
        fill=PatternFill("solid", fgColor="FFDDDD")))
    chart = BarChart()
    chart.title = "Amount by SKU"
    chart.add_data(Reference(sheet, min_col=3, min_row=1, max_row=4), titles_from_data=True)
    chart.set_categories(Reference(sheet, min_col=1, min_row=2, max_row=4))
    sheet.add_chart(chart, "H5")
    sheet.add_image(Image(str(image_path)), "H20")
    book.defined_names.add(DefinedName("Amounts", attr_text="'Sales'!$C$2:$C$4"))
    hidden = book.create_sheet("Hidden support")
    hidden.append(["reference", "value"])
    hidden.append(["keep", 42])
    hidden.sheet_state = "hidden"
    review = book.create_sheet("Review, 2026")
    review.append(["Last, First", "Status"])
    review.append(["Smith, Alice", "pending"])
    path = directory / "representative-advanced.xlsx"
    book.save(path)
    book.close()
    return path
