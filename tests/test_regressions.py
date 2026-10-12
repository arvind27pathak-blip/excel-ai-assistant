import asyncio
import uuid
import unittest
from pathlib import Path
from unittest.mock import Mock, patch

import pandas as pd
from openpyxl import Workbook, load_workbook

from app.services.data_manager import DataManager
from app.services.api_manager import APIManager
from app.services.ollama_api_manager import OllamaAPIManager
from app.ui.components.batch_processor import BatchProcessor
from app.ui.main_window import ExcelAIAssistantApp


class DataTests(unittest.TestCase):
    def setUp(self):
        self.directory = Path(__file__).resolve().parent / ('fixtures-' + uuid.uuid4().hex)
        self.directory.mkdir()
        def cleanup():
            for item in self.directory.iterdir():
                item.unlink()
            self.directory.rmdir()
        self.addCleanup(cleanup)
        self.source = self.directory / 'source.csv'
        self.source.write_text('name,amount\nAlice,12\nBob,23\n', encoding='utf-8')
        self.manager = DataManager()
        self.assertTrue(self.manager.load_file(self.source)[0])

    def test_failed_load_keeps_unsaved_data(self):
        self.manager.update_cell(0, 'name', 'Edited')
        empty = self.directory / 'empty.csv'
        empty.write_text('header\n', encoding='utf-8')
        self.assertFalse(self.manager.load_file(empty)[0])
        self.assertEqual(self.manager.get_cell_value(0, 'name'), 'Edited')
        self.assertEqual(self.manager.file_path, self.source)
        self.assertTrue(self.manager.modified)

    def test_rejected_save_keeps_dirty_state(self):
        self.manager.update_cell(0, 'name', 'Edited')
        target = self.directory / 'output.xyz'
        self.assertFalse(self.manager.save_file(target)[0])
        self.assertFalse(target.exists())
        self.assertTrue(self.manager.modified)
        self.assertEqual(self.manager.file_path, self.source)

    def test_save_failure_does_not_damage_source(self):
        before = self.source.read_bytes()
        self.manager.update_cell(0, 'name', 'Edited')
        with patch.object(pd.DataFrame, 'to_csv', side_effect=OSError('disk full')):
            self.assertFalse(self.manager.save_file()[0])
        self.assertEqual(before, self.source.read_bytes())
        self.assertTrue(self.manager.modified)
        self.assertEqual(list(self.directory.iterdir()), [self.source])

    def test_extensionless_save_and_metadata(self):
        self.assertTrue(self.manager.save_file(self.directory / 'output')[0])
        self.assertTrue((self.directory / 'output.xlsx').exists())
        self.assertEqual(self.manager.file_type, 'excel')

    def test_numeric_cell_accepts_ai_text(self):
        self.assertTrue(self.manager.update_cell(0, 'amount', 'twelve'))
        self.assertEqual(self.manager.get_cell_value(0, 'amount'), 'twelve')

    def test_numeric_headers(self):
        path = self.directory / 'numeric.xlsx'
        pd.DataFrame([[1, 2]], columns=[100, 200]).to_excel(path, index=False)
        self.assertTrue(self.manager.load_file(path)[0])
        self.assertEqual(list(self.manager.df.columns), ['100', '200'])

    def test_duplicate_normalized_headers_rejected(self):
        path = self.directory / 'duplicate.csv'
        path.write_text('name, name\nAlice,Bob\n', encoding='utf-8')
        self.assertFalse(self.manager.load_file(path)[0])
        self.assertEqual(self.manager.file_path, self.source)

    def test_excel_preserves_other_sheets_formulas_and_formatting(self):
        path = self.directory / 'workbook.xlsx'
        book = Workbook()
        sheet = book.active
        sheet.append(['name', 'formula'])
        sheet.append(['Alice', '=1+2'])
        sheet['A2'].number_format = '@'
        book.create_sheet('Other')['A1'] = 'Keep me'
        book.save(path)
        book.close()
        self.assertTrue(self.manager.load_file(path)[0])
        self.assertTrue(self.manager.update_cell(0, 'name', 'Edited'))
        self.assertTrue(self.manager.save_file()[0])
        book = load_workbook(path)
        self.addCleanup(book.close)
        self.assertEqual(book.active['A2'].value, 'Edited')
        self.assertEqual(book.active['A2'].number_format, '@')
        self.assertEqual(book.active['B2'].value, '=1+2')
        self.assertEqual(book['Other']['A1'].value, 'Keep me')


class BatchTests(unittest.IsolatedAsyncioTestCase):
    async def test_cancel_retains_current_batch_results(self):
        processor = BatchProcessor(Mock(), Mock(), asyncio.get_running_loop())
        processor._update_progress = Mock()
        processor._update_batch_progress = Mock()
        processor._call_completion = Mock()
        async def process(*args):
            return True, 'finished', None
        processor._process_single_cell_async = process
        cells = [{'row': i, 'col': 'name', 'content': 'old'} for i in range(3)]
        task = asyncio.create_task(processor._process_batches_async(
            cells, '', '', 10, 0.3, 150, None, Mock()))
        await asyncio.sleep(0.01)
        task.cancel()
        await task
        results, successes, errors = processor._call_completion.call_args.args[1:]
        self.assertEqual(len(results), 1)
        self.assertEqual(successes, 1)
        self.assertEqual(errors, 0)


class IntegrationTests(unittest.TestCase):
    def test_ollama_uses_selected_model_at_startup(self):
        manager = APIManager(model='custom-model', api_type='ollama')
        self.assertEqual(manager.ollama_manager.model, 'custom-model')

    def test_stream_error_does_not_become_successful_empty_cell(self):
        response = Mock(status_code=200)
        response.iter_lines.return_value = [b'{"error":"model unavailable"}']
        with patch('app.services.ollama_api_manager.requests.post', return_value=response):
            success, result, error = OllamaAPIManager().process_single_cell('old', '', '')
        self.assertFalse(success)
        self.assertIsNone(result)
        self.assertIn('model unavailable', error)
        response.close.assert_called_once()

    def test_incomplete_stream_is_rejected(self):
        response = Mock(status_code=200)
        response.iter_lines.return_value = [b'{"response":"partial"}']
        with patch('app.services.ollama_api_manager.requests.post', return_value=response):
            self.assertFalse(OllamaAPIManager().process_single_cell('old', '', '')[0])
        response.close.assert_called_once()

    def test_closed_progress_window_still_stages_results_for_review(self):
        app = ExcelAIAssistantApp.__new__(ExcelAIAssistantApp)
        app.data_manager = DataManager()
        app.data_manager.df = pd.DataFrame({'name': ['old']})
        app.data_treeview = Mock()
        app.results_treeview = Mock()
        app.batch_processor = Mock(processing_cancelled=True)
        app.save_after_batch_var = Mock()
        app.save_after_batch_var.get.return_value = False
        app.progress_window = Mock()
        app.progress_window.winfo_exists.return_value = False
        app._run_document_id = app.data_manager.document_id
        app._show_result_review = Mock()
        app.log = Mock()
        app.status_bar = Mock()
        app._processing_completed([{'row': 0, 'col': 'name', 'result': 'new', 'success': True}], 1, 0)
        self.assertEqual(app.data_manager.get_cell_value(0, 'name'), 'old')
        self.assertFalse(app.data_manager.modified)
        self.assertEqual(len(app._pending_results), 1)
        app._show_result_review.assert_called_once()

    def test_auto_save_failure_retains_completed_changes(self):
        app = ExcelAIAssistantApp.__new__(ExcelAIAssistantApp)
        app.data_manager = DataManager()
        app.data_manager.df = pd.DataFrame({'name': ['old']})
        app.data_manager.save_file = Mock(return_value=(False, 'disk full'))
        app.data_treeview = Mock()
        app.results_treeview = Mock()
        app.batch_processor = Mock(processing_cancelled=False)
        app.save_after_batch_var = Mock()
        app.save_after_batch_var.get.return_value = True
        app.progress_window = None
        app._run_document_id = app.data_manager.document_id
        app._run_autosave = True
        app.log = Mock()
        app.status_bar = Mock()
        with patch('app.ui.main_window.messagebox.showerror') as alert:
            app._apply_reviewed_results([{'row': 0, 'col': 'name', 'result': 'new', 'success': True}])
        self.assertEqual(app.data_manager.get_cell_value(0, 'name'), 'new')
        self.assertTrue(app.data_manager.modified)
        alert.assert_called_once()
        self.assertIn('save failed', app.status_bar.set_status.call_args.args[0].lower())

    def test_close_does_not_exit_if_save_fails(self):
        app = ExcelAIAssistantApp.__new__(ExcelAIAssistantApp)
        app.data_manager = Mock(modified=True, file_path='source.csv')
        app.data_manager.save_file.return_value = (False, 'Permission denied')
        app.root = Mock()
        with patch('app.ui.main_window.messagebox.askyesnocancel', return_value=True), patch('app.ui.main_window.messagebox.showerror'):
            app._on_close()
        app.root.destroy.assert_not_called()


class TreeviewTests(unittest.TestCase):
    def setUp(self):
        import tkinter as tk
        from app.ui.components.treeview import DataTreeview
        try:
            self.root = tk.Tk()
        except tk.TclError as error:
            self.skipTest(f"Tk unavailable: {error}")
        self.root.withdraw()
        self.addCleanup(self.root.destroy)
        self.view = DataTreeview(self.root)
        self.manager = DataManager()
        self.manager.df = pd.DataFrame({'name': ['z', 'a', '(special)']})
        self.view.set_data(self.manager.df)
        self.view.set_on_cell_edit_callback(self.manager.update_cell)

    def test_sorted_edit_updates_source_and_dirty_state(self):
        self.view._sort_by_column('name')
        self.view._save_cell_edit(Mock(), 1, 'name', 'edited')
        self.assertEqual(self.manager.get_cell_value(1, 'name'), 'edited')
        self.assertTrue(self.manager.modified)

    def test_filter_accepts_literal_parenthesis_and_edits_original_row(self):
        self.view.filter_data('(')
        self.assertEqual(list(self.view.df.index), [2])
        self.view._save_cell_edit(Mock(), 2, 'name', 'edited')
        self.assertEqual(self.manager.get_cell_value(2, 'name'), 'edited')

    def test_copy_filtered_row_uses_original_row_identity(self):
        self.view.filter_data('special')
        self.view._copy_row(2)
        self.assertEqual(self.root.clipboard_get(), '(special)')


if __name__ == '__main__':
    unittest.main()
